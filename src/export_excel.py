from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MODEL_DIR = PROJECT_ROOT / "data" / "model"
MARTS_DIR = PROJECT_ROOT / "data" / "marts"
REPORTS_DIR = PROJECT_ROOT / "reports"


TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build an Excel workbook from modeled CSV tables.")
    parser.add_argument(
        "--output",
        default=str(REPORTS_DIR / "retail_electronics_model.xlsx"),
        help="Output .xlsx path.",
    )
    parser.add_argument(
        "--fact-row-limit",
        type=int,
        default=100_000,
        help="Maximum fact rows to include in the Excel workbook for file-size control.",
    )
    return parser.parse_args()


def read_csv(path: Path, limit: int | None = None) -> tuple[list[str], list[list[str]]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        rows = []
        for index, row in enumerate(reader):
            if limit is not None and index >= limit:
                break
            rows.append(row)
    return headers, rows


def add_table_sheet(
    wb: Workbook,
    sheet_name: str,
    csv_path: Path,
    table_name: str,
    row_limit: int | None = None,
) -> None:
    headers, rows = read_csv(csv_path, limit=row_limit)
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(row)

    style_sheet(ws)
    if ws.max_row >= 2 and ws.max_column >= 1:
        table_ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = TABLE_STYLE
        ws.add_table(table)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2DD")
    border = Border(bottom=thin)

    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    for column in ws.columns:
        header = str(column[0].value or "")
        width = max(len(header) + 2, 12)
        for cell in column[1:80]:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(width, 42)


def add_readme_sheet(wb: Workbook, fact_limit: int) -> None:
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "Retail Electronics Analytics Model"
    ws["A1"].font = Font(size=20, bold=True, color="111313")
    ws["A3"] = "Purpose"
    ws["B3"] = "Practice data modeling from raw Amazon Electronics review JSONL into fact/dimension tables, marts, and dashboard-ready outputs."
    ws["A5"] = "Grain"
    ws["B5"] = "fact_reviews: one row per review. Product dimension grain: one row per ASIN. Month dimension grain: one row per review_month."
    ws["A7"] = "Included fact rows"
    ws["B7"] = fact_limit
    ws["A9"] = "Recommended workflow"
    ws["B9"] = "Use fact/dim sheets to understand the model, marts for dashboarding, and Dashboard sheet for executive-facing summary."
    ws["A11"] = "Source"
    ws["B11"] = "Electronics_5.json local raw extract. Raw file is intentionally excluded from Git."

    for row in [3, 5, 7, 9, 11]:
        ws[f"A{row}"].font = Font(bold=True, color="0F766E")
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110


def add_dashboard_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Dashboard", 1)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Retail Electronics Review Dashboard"
    ws["A1"].font = Font(size=22, bold=True, color="111313")
    ws["A2"] = "Executive snapshot generated from modeled review data and dashboard marts."
    ws["A2"].font = Font(color="646866")

    metrics = read_key_value(MARTS_DIR / "mart_data_quality_summary.csv")
    product_headers, product_rows = read_csv(MARTS_DIR / "mart_product_performance.csv", limit=8)
    rating_headers, rating_rows = read_csv(MARTS_DIR / "mart_rating_distribution.csv")
    monthly_headers, monthly_rows = read_csv(MARTS_DIR / "mart_monthly_review_trends.csv")

    kpis = [
        ("Processed reviews", metrics.get("rows_valid", "")),
        ("Unique products", metrics.get("unique_products", "")),
        ("Review text coverage", f"{metrics.get('pct_rows_with_review_text', '')}%"),
        ("Helpful vote coverage", f"{metrics.get('pct_rows_with_helpful_votes', '')}%"),
    ]

    start_cols = ["A", "C", "E", "G"]
    for index, (label, value) in enumerate(kpis):
        col = start_cols[index]
        ws[f"{col}4"] = label
        ws[f"{col}5"] = value
        ws[f"{col}4"].font = Font(bold=True, color="646866")
        ws[f"{col}5"].font = Font(size=18, bold=True, color="111313")
        ws.merge_cells(f"{col}4:{get_column_letter(ord(col) - ord('A') + 2)}4")
        ws.merge_cells(f"{col}5:{get_column_letter(ord(col) - ord('A') + 2)}5")

    ws["A8"] = "Top products by review volume"
    ws["A8"].font = Font(size=14, bold=True)
    top_cols = ["asin", "review_count", "avg_rating", "quality_flag"]
    for col_index, header in enumerate(top_cols, 1):
        ws.cell(row=9, column=col_index, value=header)
    header_positions = {header: product_headers.index(header) for header in top_cols}
    for row_index, row in enumerate(product_rows, 10):
        for col_index, header in enumerate(top_cols, 1):
            ws.cell(row=row_index, column=col_index, value=row[header_positions[header]])

    ws["F8"] = "Rating distribution"
    ws["F8"].font = Font(size=14, bold=True)
    for col_index, header in enumerate(rating_headers, 6):
        ws.cell(row=9, column=col_index, value=header)
    for row_index, row in enumerate(rating_rows, 10):
        for col_index, value in enumerate(row, 6):
            ws.cell(row=row_index, column=col_index, value=value)

    ws["A21"] = "Monthly trend source"
    ws["A21"].font = Font(size=14, bold=True)
    trend_rows = monthly_rows[-24:]
    for col_index, header in enumerate(monthly_headers, 1):
        ws.cell(row=22, column=col_index, value=header)
    for row_index, row in enumerate(trend_rows, 23):
        for col_index, value in enumerate(row, 1):
            ws.cell(row=row_index, column=col_index, value=value)

    style_sheet(ws)
    add_dashboard_charts(ws, len(rating_rows), len(trend_rows))


def read_key_value(path: Path) -> dict[str, str]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return {row["metric"]: row["value"] for row in reader}


def add_dashboard_charts(ws, rating_count: int, trend_count: int) -> None:
    bar = BarChart()
    bar.title = "Rating Distribution"
    bar.y_axis.title = "Reviews"
    bar.x_axis.title = "Rating"
    rating_data = Reference(ws, min_col=7, min_row=9, max_row=9 + rating_count)
    rating_categories = Reference(ws, min_col=6, min_row=10, max_row=9 + rating_count)
    bar.add_data(rating_data, titles_from_data=True)
    bar.set_categories(rating_categories)
    bar.height = 7
    bar.width = 13
    ws.add_chart(bar, "F17")

    line = LineChart()
    line.title = "Monthly Review Volume"
    line.y_axis.title = "Reviews"
    line.x_axis.title = "Month"
    line_data = Reference(ws, min_col=2, min_row=22, max_row=22 + trend_count)
    line_categories = Reference(ws, min_col=1, min_row=23, max_row=22 + trend_count)
    line.add_data(line_data, titles_from_data=True)
    line.set_categories(line_categories)
    line.height = 7
    line.width = 18
    ws.add_chart(line, "A49")


def add_model_relationship_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Model Relationships")
    rows = [
        ["Table", "Grain", "Primary key", "Joins"],
        ["fact_reviews", "One row per review", "review_id", "asin -> dim_products, review_month -> dim_review_months, rating_id -> dim_ratings"],
        ["dim_products", "One row per ASIN", "asin", "fact_reviews.asin"],
        ["dim_review_months", "One row per review month", "review_month", "fact_reviews.review_month"],
        ["dim_ratings", "One row per rating value", "rating_id", "fact_reviews.rating_id"],
        ["mart_product_performance", "One row per ASIN", "asin", "Dashboard source"],
        ["mart_monthly_review_trends", "One row per review month", "review_month", "Dashboard source"],
        ["mart_rating_distribution", "One row per rating", "rating_id/rating", "Dashboard source"],
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    ws.column_dimensions["D"].width = 90


def build_workbook(output_path: Path, fact_limit: int) -> None:
    wb = Workbook()
    add_readme_sheet(wb, fact_limit)
    add_dashboard_sheet(wb)
    add_model_relationship_sheet(wb)

    add_table_sheet(wb, "fact_reviews", MODEL_DIR / "fact_reviews.csv", "FactReviews", row_limit=fact_limit)
    add_table_sheet(wb, "dim_products", MODEL_DIR / "dim_products.csv", "DimProducts")
    add_table_sheet(wb, "dim_review_months", MODEL_DIR / "dim_review_months.csv", "DimReviewMonths")
    add_table_sheet(wb, "dim_ratings", MODEL_DIR / "dim_ratings.csv", "DimRatings")
    add_table_sheet(wb, "mart_product_performance", MARTS_DIR / "mart_product_performance.csv", "MartProductPerformance")
    add_table_sheet(wb, "mart_monthly_trends", MARTS_DIR / "mart_monthly_review_trends.csv", "MartMonthlyTrends")
    add_table_sheet(wb, "mart_rating_distribution", MARTS_DIR / "mart_rating_distribution.csv", "MartRatingDistribution")
    add_table_sheet(wb, "mart_data_quality", MARTS_DIR / "mart_data_quality_summary.csv", "MartDataQuality")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    output_path = Path(args.output).expanduser().resolve()
    build_workbook(output_path, args.fact_row_limit)
    print(f"Wrote Excel workbook to {output_path}")


if __name__ == "__main__":
    main()

